from __future__ import annotations

import argparse
import os
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_realdisp_manuscript_artifacts as base  # noqa: E402


EAST_ASIA_BODY = "SimSun"
EAST_ASIA_HEADING = "SimHei"
LATIN_FONT = "Times New Roman"


def set_mixed_font(
    run,
    east_asia: str,
    size: float,
    *,
    bold: bool | None = None,
) -> None:
    run.font.name = LATIN_FONT
    rfonts = run._element.get_or_add_rPr().rFonts
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = base.INK
    if bold is not None:
        run.bold = bold


def set_style_font(style, east_asia: str, latin: str, size: float) -> None:
    style.font.name = latin
    rfonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    style.font.size = Pt(size)


def add_chinese_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = paragraph.add_run("第 ")
    set_mixed_font(prefix, EAST_ASIA_BODY, 8.5)
    field_run = paragraph.add_run()
    set_mixed_font(field_run, EAST_ASIA_BODY, 8.5)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.extend((begin, instruction, separate, value, end))
    suffix = paragraph.add_run(" 页")
    set_mixed_font(suffix, EAST_ASIA_BODY, 8.5)


def apply_chinese_styles(path: Path) -> None:
    document = Document(path)
    set_style_font(document.styles["Normal"], EAST_ASIA_BODY, LATIN_FONT, 10.5)
    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        set_style_font(document.styles[name], EAST_ASIA_HEADING, LATIN_FONT, size)
    if "Caption" in document.styles:
        set_style_font(document.styles["Caption"], EAST_ASIA_BODY, LATIN_FONT, 8.5)

    references_started = False
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if index == 0:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(7)
            for run in paragraph.runs:
                set_mixed_font(run, EAST_ASIA_HEADING, 18, bold=True)
            continue
        if text.startswith("匿名作者"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(2)
            for run in paragraph.runs:
                set_mixed_font(run, EAST_ASIA_BODY, 10.5)
            continue
        if re.match(r"^表\s*[1-5][.．]", text):
            paragraph.style = "Caption"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                set_mixed_font(run, EAST_ASIA_BODY, 8.5, bold=True)
            continue
        if text == "参考文献":
            references_started = True
        elif references_started:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.0

        east_asia = EAST_ASIA_HEADING if style_name.startswith("Heading") else EAST_ASIA_BODY
        size = 10.5
        if style_name == "Heading 1":
            size = 14
        elif style_name == "Heading 2":
            size = 12
        elif style_name == "Heading 3":
            size = 11
        elif style_name == "Caption":
            size = 8.5
        elif references_started:
            size = 9.0
        for run in paragraph.runs:
            set_mixed_font(run, east_asia, size)

    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_mixed_font(run, EAST_ASIA_BODY, 8.0, bold=row_index == 0)

    section = document.sections[0]
    header_paragraph = section.header.paragraphs[0]
    header_paragraph.text = "FMC-PC｜因子化边缘校准预测一致性"
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_paragraph.paragraph_format.space_after = Pt(0)
    for run in header_paragraph.runs:
        set_mixed_font(run, EAST_ASIA_BODY, 8.5)
        run.font.color.rgb = base.MUTED

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.clear()
    add_chinese_page_number(footer_paragraph)

    core = document.core_properties
    core.title = "自主佩戴偏移鲁棒活动识别：因子化边缘校准预测一致性"
    core.author = "匿名作者"
    core.subject = "自主佩戴偏移条件下的可穿戴人体活动识别"
    core.keywords = "可穿戴人体活动识别；传感器位移；旋转增强；一致性正则化"
    document.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the Chinese FMC-PC manuscript DOCX.")
    parser.add_argument("--manuscript-dir", type=Path, required=True)
    parser.add_argument("--build-dir", type=Path, required=True)
    args = parser.parse_args()

    manuscript_dir = args.manuscript_dir.resolve()
    build_dir = args.build_dir.resolve()
    build_dir.mkdir(parents=True, exist_ok=True)
    source = manuscript_dir / "manuscript_zh.md"
    temporary = build_dir / "manuscript_zh_docx.md"
    output = manuscript_dir / "FMC_PC_manuscript_zh.docx"
    base.prepare_markdown(source, temporary)
    base.run(
        [
            str(base.PANDOC),
            str(temporary),
            "--citeproc",
            "--csl",
            str(manuscript_dir / "ieee.csl"),
            "--bibliography",
            str(manuscript_dir / "references.bib"),
            "--resource-path",
            os.pathsep.join(
                (str(manuscript_dir), str(manuscript_dir.parent / "figures"))
            ),
            "-o",
            str(output),
        ],
        manuscript_dir,
    )
    base.style_document(output, supplement=False)
    apply_chinese_styles(output)
    print(f"Chinese manuscript built: {output}")


if __name__ == "__main__":
    main()
