from __future__ import annotations

import argparse
import os
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# The DOCX resolves the narrative_proposal preset into a restrained scientific
# manuscript override: Letter portrait, 1-in margins, Times New Roman 10.5 pt,
# justified 1.10-line body, black 14/12/11-pt headings, fixed-width tables,
# light-gray headers, inline figures, and quiet running page furniture.
PANDOC = Path(r"D:\edg\anaconda\Library\bin\pandoc.exe")
XELATEX = Path(
    r"C:\Users\20921\AppData\Local\Programs\MiKTeX\miktex\bin\x64\xelatex.exe"
)
INK = RGBColor(0x18, 0x18, 0x18)
MUTED = RGBColor(0x55, 0x55, 0x55)
LIGHT_GRAY = "E9EDF2"


def run(command: list[str], cwd: Path) -> None:
    print("RUN", " ".join(command), flush=True)
    subprocess.run(command, cwd=cwd, check=True)


def set_font(run, name: str, size: float, *, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = INK
    if bold is not None:
        run.bold = bold


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row_index, row in enumerate(table.rows):
        for column, (cell, width) in enumerate(zip(row.cells, widths_dxa, strict=True)):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if column < 2 else WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in paragraph.runs:
                    set_font(run, "Times New Roman", 8.0, bold=row_index == 0)
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
            for cell in row.cells:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), LIGHT_GRAY)
                cell._tc.get_or_add_tcPr().append(shading)


def table_pattern(table, usable_dxa: int) -> list[int]:
    columns = len(table.columns)
    if usable_dxa == 9360:
        patterns = {
            2: [2880, 6480],
            5: [2300, 1765, 1765, 1765, 1765],
            6: [1150, 1700, 1600, 1600, 1500, 1810],
            7: [1800, 1000, 1250, 1800, 1000, 850, 1660],
        }
        if columns in patterns:
            return patterns[columns]
    base = usable_dxa // columns
    widths = [base] * columns
    widths[-1] += usable_dxa - sum(widths)
    return widths


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, "Times New Roman", 8.5)
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char, instr, separate, value, end))


def style_document(path: Path, *, supplement: bool = False) -> None:
    document = Document(path)
    section = document.sections[0]
    if supplement:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.0)
        section.page_height = Inches(8.5)
        margins = 0.6
        usable_dxa = 13392
    else:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        margins = 1.0
        usable_dxa = 9360
    section.top_margin = Inches(margins)
    section.bottom_margin = Inches(margins)
    section.left_margin = Inches(margins)
    section.right_margin = Inches(margins)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(9.5 if supplement else 10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, before, after in (
        ("Heading 1", 14, 12, 6),
        ("Heading 2", 12, 10, 4),
        ("Heading 3", 11, 8, 3),
    ):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = INK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Caption" in document.styles:
        caption = document.styles["Caption"]
        caption.font.name = "Times New Roman"
        caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        caption.font.size = Pt(8.5)
        caption.font.italic = False
        caption.font.color.rgb = MUTED
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(3)
        caption.paragraph_format.space_after = Pt(7)
        caption.paragraph_format.keep_together = True

    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if index == 0 and not supplement:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(7)
            for run in paragraph.runs:
                set_font(run, "Times New Roman", 18, bold=True)
        elif text.startswith("Anonymous Author(s)"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(2)
            for run in paragraph.runs:
                set_font(run, "Times New Roman", 10.5)
        elif re.match(r"^Table [1-4]\.", text):
            paragraph.style = "Caption"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                set_font(run, "Times New Roman", 8.5, bold=True)
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(2)
            for drawing in paragraph._p.xpath(".//wp:inline/wp:extent"):
                max_width = int((10.0 if supplement else 6.35) * 914400)
                max_height = int((5.7 if supplement else 8.3) * 914400)
                width = int(drawing.get("cx"))
                height = int(drawing.get("cy"))
                if width > max_width:
                    height = round(height * max_width / width)
                    width = max_width
                if height > max_height:
                    width = round(width * max_height / height)
                    height = max_height
                drawing.set("cx", str(width))
                drawing.set("cy", str(height))

    for table_index, table in enumerate(document.tables):
        set_table_widths(table, table_pattern(table, usable_dxa))
        if not supplement and table_index == 0:
            for row in table.rows[:-1]:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True

    references_started = False
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "References":
            references_started = True
            continue
        if references_started:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    header = section.header
    hp = header.paragraphs[0]
    hp.text = (
        "PD-PC | Supplementary material"
        if supplement
        else "Placement-Distribution-Calibrated Prediction Consistency"
    )
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    for run in hp.runs:
        set_font(run, "Times New Roman", 8.5)
        run.font.color.rgb = MUTED
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.clear()
    add_page_number(fp)

    core = document.core_properties
    core.title = (
        "PD-PC Supplementary Material"
        if supplement
        else "Placement-Distribution-Calibrated Prediction Consistency"
    )
    core.author = "Anonymous Author(s)"
    core.subject = "Wearable human activity recognition under Self-placement shift"
    core.keywords = "wearable HAR; sensor displacement; consistency regularization"
    document.save(path)


def prepare_markdown(source: Path, destination: Path) -> None:
    text = source.read_text(encoding="utf-8")
    text = text.replace(".pdf)", ".png)")
    text = re.sub(r"\{#fig:[^}]+ width=100%\}", "", text)
    destination.write_text(text, encoding="utf-8")


def build_docx(manuscript_dir: Path, build_dir: Path) -> None:
    csl = manuscript_dir / "ieee.csl"
    bibliography = manuscript_dir / "references.bib"
    for stem, supplement in (("manuscript", False), ("supplementary", True)):
        temp_md = build_dir / f"{stem}_docx.md"
        prepare_markdown(manuscript_dir / f"{stem}.md", temp_md)
        output = manuscript_dir / f"PD_PC_{stem}.docx"
        run(
            [
                str(PANDOC),
                str(temp_md),
                "--citeproc",
                "--csl",
                str(csl),
                "--bibliography",
                str(bibliography),
                "--resource-path",
                os.pathsep.join((str(manuscript_dir), str(manuscript_dir.parent / "figures"))),
                "-o",
                str(output),
            ],
            manuscript_dir,
        )
        style_document(output, supplement=supplement)


def build_latex_and_pdf(manuscript_dir: Path) -> None:
    latex_source = manuscript_dir / "build" / "manuscript_latex.md"
    latex_source.parent.mkdir(parents=True, exist_ok=True)
    latex_text = (manuscript_dir / "manuscript.md").read_text(encoding="utf-8")
    latex_text = re.sub(r"!\[Figure\s+\d+\.\s*", "![", latex_text)
    latex_source.write_text(latex_text, encoding="utf-8")
    common = [
        str(PANDOC),
        str(latex_source.relative_to(manuscript_dir)),
        "--standalone",
        "--citeproc",
        "--csl",
        "ieee.csl",
        "--bibliography",
        "references.bib",
        "--resource-path",
        os.pathsep.join((str(manuscript_dir), str(manuscript_dir.parent / "figures"))),
        "-V",
        "geometry:margin=0.75in",
        "-V",
        "mainfont=Times New Roman",
        "-V",
        "fontsize=10pt",
        "-V",
        "colorlinks=true",
        "-V",
        "linkcolor=black",
        "-V",
        "citecolor=black",
        "-V",
        "urlcolor=blue",
    ]
    run(common + ["-t", "latex", "-o", "PD_PC_manuscript.tex"], manuscript_dir)
    run(
        common
        + [
            "--pdf-engine",
            str(XELATEX),
            "-o",
            "PD_PC_manuscript.pdf",
        ],
        manuscript_dir,
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Build PD-PC manuscript artifacts.")
    parser.add_argument("--manuscript-dir", type=Path, required=True)
    parser.add_argument("--build-dir", type=Path, required=True)
    parser.add_argument(
        "--skip-latex-pdf",
        action="store_true",
        help="Build the DOCX files only; retain any existing LaTeX/PDF artifacts.",
    )
    args = parser.parse_args()
    manuscript_dir = args.manuscript_dir.resolve()
    build_dir = args.build_dir.resolve()
    if build_dir.exists():
        shutil.rmtree(build_dir)
    build_dir.mkdir(parents=True)
    build_docx(manuscript_dir, build_dir)
    if not args.skip_latex_pdf:
        build_latex_and_pdf(manuscript_dir)
    print("artifacts built in", manuscript_dir)


if __name__ == "__main__":
    main()
